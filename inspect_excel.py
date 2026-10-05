import openpyxl, os, sys

files = {
    'TMD': 'data/Track_Management_Department.xlsx',
    'SNT': 'data/Indian_Railway_S&T_Management.xlsx',
    'TRD': 'data/Indian_Railways_Track_Distribution_System.xlsx'
}

for dept, path in files.items():
    if not os.path.exists(path):
        print(f'MISSING: {path}')
        continue
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f'\n{"="*60}')
    print(f'DEPT: {dept}  |  FILE: {path}')
    print(f'Sheets: {wb.sheetnames}')
    print(f'{"="*60}')
    for name in wb.sheetnames:
        ws = wb[name]
        all_rows = list(ws.iter_rows(values_only=True))
        # find first non-blank row as header
        header_row = None
        data_start = 0
        for i, row in enumerate(all_rows[:15]):
            non_null = [c for c in row if c is not None and str(c).strip() != '']
            if len(non_null) >= 2:
                header_row = row
                data_start = i + 1
                break
        if header_row is None:
            print(f'  Sheet "{name}": EMPTY or no usable header')
            continue
        headers = [str(c).strip() if c is not None else '' for c in header_row]
        headers = [h for h in headers if h]
        data_rows = [r for r in all_rows[data_start:] if any(c is not None and str(c).strip() != '' for c in r)]
        print(f'\n  Sheet: "{name}"')
        print(f'  Rows: {len(data_rows)}  |  Columns: {len(headers)}')
        print(f'  Headers:')
        for h in headers:
            print(f'    [{h}]')
        # show 2 sample data rows
        if data_rows:
            print(f'  Sample row 1: {[str(v)[:40] if v is not None else "" for v in data_rows[0][:len(headers)]]}')
            if len(data_rows) > 1:
                print(f'  Sample row 2: {[str(v)[:40] if v is not None else "" for v in data_rows[1][:len(headers)]]}')
    wb.close()
print('\nDone.')
