import sys, os
print('start', flush=True)
import openpyxl
os.chdir(r'c:\Users\THISHANTH T\Desktop\PROTOTYPE')
print('loading 3dept...', flush=True)
# Use read_only=True to avoid memory issues
wb = openpyxl.load_workbook('ml/data/3dept.xlsx', read_only=True, data_only=True)
print('sheets:', wb.sheetnames[:3], flush=True)
# Read just the header row of Maintenance
ws = wb['Maintenance']
headers = []
for row in ws.iter_rows(max_row=1, values_only=True):
    headers = [c for c in row if c is not None]
    break
print('Maintenance headers[:5]:', headers[:5], flush=True)
wb.close()
print('3dept: OK', flush=True)
