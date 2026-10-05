"""Inspect remaining TRD sheets only (after Train_Stops which timed out)."""
import openpyxl, os

path = 'data/Indian_Railways_Track_Distribution_System.xlsx'
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

# Skip first 6 sheets already captured
skip = {'README', 'Zones', 'Divisions', 'Stations', 'Trains', 'Train_Stops'}

for name in wb.sheetnames:
    if name in skip:
        continue
    ws = wb[name]
    all_rows = list(ws.iter_rows(values_only=True))
    header_row = None
    data_start = 0
    for i, row in enumerate(all_rows[:15]):
        non_null = [c for c in row if c is not None and str(c).strip() != '']
        if len(non_null) >= 2:
            header_row = row
            data_start = i + 1
            break
    if header_row is None:
        print(f'  Sheet "{name}": EMPTY')
        continue
    headers = [str(c).strip() if c is not None else '' for c in header_row]
    headers = [h for h in headers if h]
    data_rows = [r for r in all_rows[data_start:] if any(c is not None and str(c).strip() != '' for c in r)]
    print(f'\nSheet: "{name}"  rows={len(data_rows)}  cols={len(headers)}')
    for h in headers:
        print(f'  [{h}]')
    if data_rows:
        print(f'  Sample: {[str(v)[:35] if v is not None else "" for v in data_rows[0][:len(headers)]]}')

wb.close()
print('\nDone.')
