import openpyxl, os, sys
os.chdir(r'c:\Users\THISHANTH T\Desktop\PROTOTYPE')
print("openpyxl version:", openpyxl.__version__)
try:
    wb = openpyxl.load_workbook('ml/data/3dept.xlsx', read_only=True, data_only=True)
    print("3dept sheets:", wb.sheetnames[:5])
    ws = wb['Maintenance']
    h = [c.value for c in next(ws.iter_rows(max_row=1))]
    print("Maintenance headers[:6]:", h[:6])
    wb.close()
    print("openpyxl READ: OK")
except Exception as e:
    import traceback; traceback.print_exc()
    print("openpyxl READ: FAIL", e)
    sys.exit(1)

# Test CSV
import csv
try:
    with open('ml/data/India_Railway_Stations_State_District_Wise (1).csv', 'r', encoding='utf-8', errors='replace') as f:
        reader = csv.DictReader(f)
        h2 = reader.fieldnames
        r1 = next(reader)
    print("CSV headers[:4]:", h2[:4])
    print("CSV row1 sample:", dict(list(r1.items())[:2]))
    print("CSV READ: OK")
except Exception as e:
    print("CSV READ: FAIL", e)

# Test sklearn predict without pandas
import joblib, numpy as np
try:
    art = joblib.load('ml/models/ai_assistant_model.joblib')
    pipeline = art['pipeline']
    result = pipeline.predict_proba(['What is TMS?'])
    idx = result[0].argmax()
    print("sklearn predict OK: intent=%s conf=%.4f" % (pipeline.classes_[idx], result[0][idx]))
except Exception as e:
    import traceback; traceback.print_exc()
    print("sklearn predict FAIL:", e)

print("ALL DONE")
